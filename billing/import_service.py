"""Excel/CSV invoice import validator and confirm service."""

from __future__ import annotations

import csv
import io
import re
from collections import OrderedDict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from billing.approval import (
    ACTION_APPROVED,
    ACTION_CREATED,
    STATUS_APPROVED,
    log_approval_action,
)
from billing.gst import calculate_gst_breakdown, resolve_client_gst_type
from billing.models import Invoice, InvoiceImportBatch, InvoiceImportItem, InvoiceLineItem
from clients.models import Client
from config.company import COMPANY
from document_generator.constants import DOC_INVOICE
from document_generator.services import generate_document_number, get_year_series, parse_document_number

TEMPLATE_HEADERS = [
    'Invoice Number',
    'Invoice Date',
    'Due Date',
    'Customer Name',
    'PO Number',
    'PO Date',
    'Billing Period From',
    'Billing Period To',
    'Service Title',
    'Item Description',
    'HSN/SAC',
    'Unit',
    'Qty',
    'Rate',
    'GST %',
    'Remarks',
]

REQUIRED_HEADERS = (
    'Customer Name',
    'Item Description',
    'Qty',
    'Rate',
    'GST %',
)

MAX_ROWS = 500
MAX_FILE_BYTES = 5 * 1024 * 1024
COMPANY_GST_RATE = Decimal(str(COMPANY['gst_rate_percent']))

_HEADER_ALIASES = {
    'invoicenumber': 'Invoice Number',
    'invoiceno': 'Invoice Number',
    'invoicenum': 'Invoice Number',
    'invoicedate': 'Invoice Date',
    'date': 'Invoice Date',
    'duedate': 'Due Date',
    'customername': 'Customer Name',
    'customer': 'Customer Name',
    'customercode': 'Customer Name',
    'clientname': 'Customer Name',
    'client': 'Customer Name',
    'ponumber': 'PO Number',
    'poso': 'PO Number',
    'posono': 'PO Number',
    'podate': 'PO Date',
    'billingperiodfrom': 'Billing Period From',
    'periodfrom': 'Billing Period From',
    'billingperiodto': 'Billing Period To',
    'periodto': 'Billing Period To',
    'servicetitle': 'Service Title',
    'service': 'Service Title',
    'itemdescription': 'Item Description',
    'description': 'Item Description',
    'particulars': 'Item Description',
    'item': 'Item Description',
    'ponumber': 'PO Number',
    'poso': 'PO Number',
    'posono': 'PO Number',
    'podate': 'PO Date',
    'billingperiodfrom': 'Billing Period From',
    'periodfrom': 'Billing Period From',
    'billingperiodto': 'Billing Period To',
    'periodto': 'Billing Period To',
    'servicetitle': 'Service Title',
    'service': 'Service Title',
    'itemdescription': 'Item Description',
    'description': 'Item Description',
    'item': 'Item Description',
    'hsnsac': 'HSN/SAC',
    'hsn': 'HSN/SAC',
    'sac': 'HSN/SAC',
    'unit': 'Unit',
    'qty': 'Qty',
    'quantity': 'Qty',
    'rate': 'Rate',
    'gst': 'GST %',
    'gstpercent': 'GST %',
    'gstpercentage': 'GST %',
    'gstrate': 'GST %',
    'remarks': 'Remarks',
    'remark': 'Remarks',
}


class ImportFileError(Exception):
    """Raised when the uploaded file cannot be parsed as an import batch."""


def build_template_workbook():
    """Return an openpyxl workbook with the official invoice import template."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Invoices'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='0F2D52')
    required_fill = PatternFill('solid', fgColor='8B1E3F')
    sample_fill = PatternFill('solid', fgColor='F4F7FB')

    for col, header in enumerate(TEMPLATE_HEADERS, 1):
        cell = ws.cell(1, col, header)
        cell.font = header_font
        cell.fill = required_fill if header in REQUIRED_HEADERS else header_fill
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.protection = Protection(locked=True)
        ws.column_dimensions[get_column_letter(col)].width = max(16, len(header) + 4)

    sample = [
        'ITSPL26270001',
        date.today().strftime('%d-%m-%Y'),
        (date.today() + timedelta(days=30)).strftime('%d-%m-%Y'),
        'Existing Customer Name',
        'PO-001',
        date.today().strftime('%d-%m-%Y'),
        '',
        '',
        'OFC Connectivity',
        'Optical fibre installation',
        COMPANY.get('default_hsn_sac', '998422'),
        'Nos',
        '1',
        '1000.00',
        str(int(COMPANY_GST_RATE)),
        '',
    ]
    for col, value in enumerate(sample, 1):
        cell = ws.cell(2, col, value)
        cell.fill = sample_fill

    gst_dv = DataValidation(
        type='decimal',
        operator='equal',
        formula1=str(int(COMPANY_GST_RATE)),
        allow_blank=True,
        showErrorMessage=True,
        errorTitle='GST %',
        error=f'GST % must be {int(COMPANY_GST_RATE)} as per company configuration.',
    )
    gst_col = TEMPLATE_HEADERS.index('GST %') + 1
    gst_dv.add(f'{get_column_letter(gst_col)}2:{get_column_letter(gst_col)}501')
    ws.add_data_validation(gst_dv)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(TEMPLATE_HEADERS))}1'

    instructions = wb.create_sheet('Instructions')
    lines = [
        'IOMS Invoice Import Template',
        '',
        '1. Use the Invoices sheet. Do not rename header cells.',
        '2. Customer Name must already exist in IOMS (exact name, case-insensitive). Customers are never auto-created.',
        f'3. GST % must be {int(COMPANY_GST_RATE)} (company rate). Other rates are rejected.',
        '4. Rows with the same Invoice Number are grouped into one invoice (multiple line items).',
        '5. Blank Invoice Number rows are grouped by Customer Name + PO Number + dates, and a number is auto-generated on confirm.',
        '6. Manual invoice numbers from Excel are kept as-is. Duplicate numbers already in IOMS are marked Already exists.',
        '7. Maximum 500 data rows and 5 MB per file. Formats: .xlsx or .csv.',
        '8. Review the preview, then Confirm. Confirm is all-or-nothing: a failure rolls back every invoice in the batch.',
        '9. Confirmed invoices are created as Approved (PDF available) with source = Import and no order link.',
        '',
        'Required columns (maroon header): Customer Name, Item Description, Qty, Rate, GST %.',
        'Date formats: DD-MM-YYYY, DD/MM/YYYY, YYYY-MM-DD.',
    ]
    instructions.column_dimensions['A'].width = 120
    title_font = Font(bold=True, size=14, color='0F2D52')
    for i, line in enumerate(lines, 1):
        cell = instructions.cell(i, 1, line)
        if i == 1:
            cell.font = title_font
    return wb


def parse_and_validate(uploaded_file, filename, user, title=''):
    """Parse an uploaded xlsx/csv file and persist a validated InvoiceImportBatch."""
    if uploaded_file is None:
        raise ImportFileError('Please choose an Excel or CSV file to upload.')
    filename = (filename or getattr(uploaded_file, 'name', '') or '').strip()
    if not filename:
        raise ImportFileError('The uploaded file has no name.')

    ext = _file_extension(filename)
    if ext not in ('.xlsx', '.csv'):
        raise ImportFileError('Only .xlsx and .csv files are supported.')

    content = _read_upload_bytes(uploaded_file)
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError('File is larger than 5 MB.')
    if not content.strip():
        raise ImportFileError('The uploaded file is empty.')

    raw_rows = _load_rows(content, ext)
    if not raw_rows:
        raise ImportFileError('The file has no header row.')

    header_map = _map_headers(raw_rows[0])
    missing = [h for h in REQUIRED_HEADERS if h not in header_map.values()]
    if missing:
        raise ImportFileError('Missing required column(s): ' + ', '.join(missing))

    data_rows = raw_rows[1:]
    if len(data_rows) > MAX_ROWS:
        raise ImportFileError(f'Maximum {MAX_ROWS} data rows allowed per file.')

    parsed_rows = []
    for offset, raw in enumerate(data_rows, start=2):
        if _row_is_empty(raw):
            continue
        parsed_rows.append(_parse_source_row(raw, header_map, offset))

    if not parsed_rows:
        raise ImportFileError('The file has no data rows.')

    groups = _group_rows(parsed_rows)
    existing_numbers = _existing_invoice_numbers()

    batch = InvoiceImportBatch(
        title=(title or '').strip() or filename,
        original_filename=filename,
        uploaded_by=user,
        status=InvoiceImportBatch.STATUS_VALIDATED,
        total_rows=len(parsed_rows),
        invoice_count=len(groups),
    )
    safe_name = filename.replace('/', '_')
    batch.file.save(safe_name, ContentFile(content), save=False)
    batch.save()

    valid_count = warning_count = error_count = skipped_count = 0
    for sort_order, (grouping_key, rows) in enumerate(groups.items()):
        item = _build_import_item(batch, grouping_key, rows, existing_numbers, sort_order)
        item.save()
        if item.status == InvoiceImportItem.STATUS_VALID:
            valid_count += 1
        elif item.status == InvoiceImportItem.STATUS_WARNING:
            warning_count += 1
        elif item.status == InvoiceImportItem.STATUS_ALREADY_EXISTS:
            skipped_count += 1
        else:
            error_count += 1

    batch.valid_count = valid_count
    batch.warning_count = warning_count
    batch.error_count = error_count
    batch.skipped_count = skipped_count
    batch.save(update_fields=[
        'valid_count', 'warning_count', 'error_count', 'skipped_count',
    ])
    return batch


def creatable_items(batch):
    return batch.items.filter(status__in=[
        InvoiceImportItem.STATUS_VALID,
        InvoiceImportItem.STATUS_WARNING,
    ])


def batch_totals(batch, statuses=None):
    items = batch.items.all()
    if statuses is not None:
        items = items.filter(status__in=list(statuses))
    taxable = Decimal('0.00')
    gst = Decimal('0.00')
    total = Decimal('0.00')
    count = 0
    for item in items:
        payload = item.payload or {}
        taxable += _decimal(payload.get('taxable'))
        gst += _decimal(payload.get('gst'))
        total += _decimal(payload.get('total'))
        count += 1
    return {
        'taxable': taxable.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
        'gst': gst.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
        'total': total.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
        'count': count,
    }


def confirm_import_batch(batch, user):
    """Create invoices for a validated batch. Failure rolls back every invoice."""
    if batch.status != InvoiceImportBatch.STATUS_VALIDATED:
        return {
            'ok': False,
            'message': 'This import batch cannot be confirmed.',
            'created': 0,
        }

    items = list(creatable_items(batch).select_related('client'))
    if not items:
        InvoiceImportBatch.objects.filter(pk=batch.pk).update(
            confirm_message='No valid invoices to import.',
        )
        return {
            'ok': False,
            'message': 'No valid invoices to import.',
            'created': 0,
        }

    try:
        with transaction.atomic():
            created_invoices = []
            for item in items:
                invoice = _create_invoice_from_item(item, user, batch)
                created_invoices.append(invoice)
                item.status = InvoiceImportItem.STATUS_CREATED
                item.created_invoice = invoice
                item.save(update_fields=['status', 'created_invoice'])

            message = (
                f'{len(created_invoices)} invoice(s) created from import '
                f'"{batch.title or batch.original_filename}".'
            )
            batch.status = InvoiceImportBatch.STATUS_CONFIRMED
            batch.created_count = len(created_invoices)
            batch.failed_count = 0
            batch.confirm_message = message
            batch.save(update_fields=[
                'status', 'created_count', 'failed_count', 'confirm_message',
            ])
            created = len(created_invoices)
    except Exception as exc:
        InvoiceImportBatch.objects.filter(pk=batch.pk).update(
            status=InvoiceImportBatch.STATUS_FAILED,
            failed_count=len(items),
            created_count=0,
            confirm_message=f'Import failed and was rolled back: {exc}',
        )
        batch.refresh_from_db()
        return {
            'ok': False,
            'message': batch.confirm_message,
            'created': 0,
        }

    batch.refresh_from_db()
    return {'ok': True, 'message': batch.confirm_message, 'created': created}


def _file_extension(filename):
    name = filename.lower().strip()
    if name.endswith('.xlsx'):
        return '.xlsx'
    if name.endswith('.csv'):
        return '.csv'
    if '.' in name:
        return '.' + name.rsplit('.', 1)[-1]
    return ''


def _read_upload_bytes(uploaded_file):
    if hasattr(uploaded_file, 'seek'):
        try:
            uploaded_file.seek(0)
        except Exception:
            pass
    reported = getattr(uploaded_file, 'size', None)
    if reported is not None:
        try:
            if int(reported) > MAX_FILE_BYTES:
                raise ImportFileError('File is larger than 5 MB.')
        except (TypeError, ValueError):
            pass
    content = uploaded_file.read()
    if isinstance(content, str):
        content = content.encode('utf-8')
    return content or b''


def _load_rows(content, ext):
    if ext == '.csv':
        text = content.decode('utf-8-sig', errors='replace')
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        return [list(row) for row in reader]

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = workbook.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row) if row else [])
        return rows
    finally:
        workbook.close()


def _canonical_header(raw):
    key = re.sub(r'[^a-z0-9]+', '', str(raw or '').strip().lower())
    if not key:
        return None
    if key in _HEADER_ALIASES:
        return _HEADER_ALIASES[key]
    for header in TEMPLATE_HEADERS:
        if re.sub(r'[^a-z0-9]+', '', header.lower()) == key:
            return header
    return None


def _map_headers(header_row):
    mapping = {}
    for idx, raw in enumerate(header_row or []):
        canonical = _canonical_header(raw)
        if canonical and canonical not in mapping.values():
            mapping[idx] = canonical
    return mapping


def _row_is_empty(raw):
    return all(cell is None or str(cell).strip() == '' for cell in (raw or []))


def _cell_str(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return '1' if value else '0'
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, 'g')
    return str(value).strip()


def _parse_date(value, excel_serial=True):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and excel_serial:
        try:
            from openpyxl.utils.datetime import from_excel
            parsed = from_excel(value)
            if isinstance(parsed, datetime):
                return parsed.date()
            if isinstance(parsed, date):
                return parsed
        except Exception:
            pass
    text = _cell_str(value)
    if not text:
        return None
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%d.%m.%Y', '%Y/%m/%d', '%d-%b-%Y', '%d %b %Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return False


def _parse_decimal(value):
    if value is None or value == '':
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = _cell_str(value).replace(',', '').replace('%', '').strip()
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return False


def _decimal(value):
    try:
        return Decimal(str(value or 0))
    except (InvalidOperation, ValueError):
        return Decimal('0')


def _money(value):
    return _decimal(value).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def _parse_source_row(raw, header_map, row_number):
    values = {}
    for idx, header in header_map.items():
        values[header] = raw[idx] if idx < len(raw) else None

    invoice_date = _parse_date(values.get('Invoice Date'))
    due_date = _parse_date(values.get('Due Date'))
    po_date = _parse_date(values.get('PO Date'))
    period_from = _parse_date(values.get('Billing Period From'))
    period_to = _parse_date(values.get('Billing Period To'))
    qty = _parse_decimal(values.get('Qty'))
    rate = _parse_decimal(values.get('Rate'))
    gst_percent = _parse_decimal(values.get('GST %'))

    return {
        'row_number': row_number,
        'invoice_number': _cell_str(values.get('Invoice Number')),
        'invoice_date': invoice_date,
        'due_date': due_date,
        'customer_name': _cell_str(values.get('Customer Name')),
        'po_number': _cell_str(values.get('PO Number')),
        'po_date': po_date,
        'billing_period_from': period_from,
        'billing_period_to': period_to,
        'service_title': _cell_str(values.get('Service Title')),
        'description': _cell_str(values.get('Item Description')),
        'hsn_sac': _cell_str(values.get('HSN/SAC')),
        'unit': _cell_str(values.get('Unit')),
        'qty': qty,
        'rate': rate,
        'gst_percent': gst_percent,
        'remarks': _cell_str(values.get('Remarks')),
    }


def _group_rows(parsed_rows):
    groups = OrderedDict()
    for row in parsed_rows:
        number = (row['invoice_number'] or '').strip()
        if number:
            key = 'NO:' + number.upper()
        else:
            key = 'BLANK:' + '|'.join([
                (row['customer_name'] or '').strip().upper(),
                (row['po_number'] or '').strip().upper(),
                _date_key(row['invoice_date']),
                _date_key(row['due_date']),
                _date_key(row['billing_period_from']),
                _date_key(row['billing_period_to']),
            ])
        groups.setdefault(key, []).append(row)
    return groups


def _date_key(value):
    if isinstance(value, date):
        return value.isoformat()
    return ''


def _existing_invoice_numbers():
    return {
        (number or '').strip().upper()
        for number in Invoice.objects.values_list('invoice_number', flat=True)
        if number
    }


def _pick_first(rows, field, default=''):
    for row in rows:
        value = row.get(field)
        if value not in (None, '', False):
            return value
    return default


def _build_import_item(batch, grouping_key, rows, existing_numbers, sort_order):
    errors = []
    warnings = []
    row_numbers = [row['row_number'] for row in rows]

    invoice_number = (_pick_first(rows, 'invoice_number') or '').strip()
    customer_name = (_pick_first(rows, 'customer_name') or '').strip()
    customers = {(row['customer_name'] or '').strip().upper() for row in rows if (row['customer_name'] or '').strip()}
    if len(customers) > 1:
        errors.append('Rows in this invoice have different customer names.')

    if not customer_name:
        errors.append('Customer Name is required.')
        client = None
    else:
        matches = list(Client.objects.filter(name__iexact=customer_name))
        if not matches:
            errors.append(
                f'Customer "{customer_name}" was not found. Customers must already exist and are never auto-created.'
            )
            client = None
        else:
            client = matches[0]
            if len(matches) > 1:
                warnings.append(f'Multiple customers named "{customer_name}"; using the first match.')

    for field, label in (
        ('invoice_date', 'Invoice Date'),
        ('due_date', 'Due Date'),
        ('po_date', 'PO Date'),
        ('billing_period_from', 'Billing Period From'),
        ('billing_period_to', 'Billing Period To'),
    ):
        if any(row.get(field) is False for row in rows):
            errors.append(f'{label} is not a valid date.')

    invoice_date = _pick_first(rows, 'invoice_date') or timezone.localdate()
    if not _pick_first(rows, 'invoice_date'):
        warnings.append('Invoice Date was blank; today was used.')
    due_date = _pick_first(rows, 'due_date') or (invoice_date + timedelta(days=30) if isinstance(invoice_date, date) else None)
    if not _pick_first(rows, 'due_date'):
        warnings.append('Due Date was blank; invoice date + 30 days was used.')

    gst_values = []
    for row in rows:
        gst_percent = row.get('gst_percent')
        if gst_percent is False:
            errors.append(f'Row {row["row_number"]}: GST % is not a valid number.')
        elif gst_percent is None:
            errors.append(f'Row {row["row_number"]}: GST % is required.')
        else:
            gst_values.append(gst_percent)
            if gst_percent != COMPANY_GST_RATE:
                errors.append(
                    f'Row {row["row_number"]}: GST % must be {int(COMPANY_GST_RATE)} '
                    f'(company rate), not {gst_percent}.'
                )

    lines = []
    taxable = Decimal('0.00')
    for row in rows:
        line_errors = []
        description = (row.get('description') or '').strip()
        if not description:
            line_errors.append(f'Row {row["row_number"]}: Item Description is required.')
        qty = row.get('qty')
        rate = row.get('rate')
        if qty is False:
            line_errors.append(f'Row {row["row_number"]}: Qty is not a valid number.')
            qty = Decimal('0')
        elif qty is None:
            line_errors.append(f'Row {row["row_number"]}: Qty is required.')
            qty = Decimal('0')
        elif qty <= 0:
            line_errors.append(f'Row {row["row_number"]}: Qty must be greater than 0.')
        if rate is False:
            line_errors.append(f'Row {row["row_number"]}: Rate is not a valid number.')
            rate = Decimal('0')
        elif rate is None:
            line_errors.append(f'Row {row["row_number"]}: Rate is required.')
            rate = Decimal('0')
        elif rate < 0:
            line_errors.append(f'Row {row["row_number"]}: Rate cannot be negative.')

        errors.extend(line_errors)
        if line_errors:
            continue

        hsn = (row.get('hsn_sac') or '').strip() or COMPANY.get('default_hsn_sac', '998422')
        if not (row.get('hsn_sac') or '').strip():
            warnings.append(f'Row {row["row_number"]}: HSN/SAC defaulted to {hsn}.')
        unit = (row.get('unit') or '').strip() or 'Nos'
        qty = Decimal(qty).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        rate = Decimal(rate).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        line_amount = (qty * rate).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        taxable += line_amount
        lines.append({
            'row': row['row_number'],
            'description': description,
            'hsn_sac': hsn,
            'unit': unit,
            'qty': str(qty),
            'rate': str(rate),
            'amount': str(line_amount),
            'remarks': row.get('remarks') or '',
        })

    if not lines and not any('Item Description' in e or 'Qty' in e or 'Rate' in e for e in errors):
        errors.append('This invoice has no line items.')

    gst_type = resolve_client_gst_type(client) if client is not None else ''
    breakdown = calculate_gst_breakdown(taxable, gst_type) if gst_type else {
        'gst': Decimal('0.00'),
        'cgst_amount': Decimal('0.00'),
        'sgst_amount': Decimal('0.00'),
        'igst_amount': Decimal('0.00'),
    }

    number_mode = Invoice.NUMBER_MODE_MANUAL if invoice_number else Invoice.NUMBER_MODE_AUTO
    if not invoice_number:
        warnings.append('Invoice Number is blank; a number will be generated on confirm from the invoice date series.')
    elif parse_document_number(invoice_number, DOC_INVOICE) is None:
        warnings.append('Invoice Number is not in the current IOMS series; it will be stored as a manual number.')

    already_exists = bool(invoice_number and invoice_number.upper() in existing_numbers)
    if already_exists:
        errors.append(f'Invoice number {invoice_number} already exists.')

    if already_exists:
        status = InvoiceImportItem.STATUS_ALREADY_EXISTS
    elif errors:
        status = InvoiceImportItem.STATUS_ERROR
    elif warnings:
        status = InvoiceImportItem.STATUS_WARNING
    else:
        status = InvoiceImportItem.STATUS_VALID

    po_date = _pick_first(rows, 'po_date')
    period_from = _pick_first(rows, 'billing_period_from')
    period_to = _pick_first(rows, 'billing_period_to')

    payload = {
        'invoice_number': invoice_number,
        'invoice_date': invoice_date.isoformat() if isinstance(invoice_date, date) else '',
        'due_date': due_date.isoformat() if isinstance(due_date, date) else '',
        'customer_name': customer_name,
        'po_number': _pick_first(rows, 'po_number') or '',
        'po_date': po_date.isoformat() if isinstance(po_date, date) else '',
        'billing_period_from': period_from.isoformat() if isinstance(period_from, date) else '',
        'billing_period_to': period_to.isoformat() if isinstance(period_to, date) else '',
        'service_title': _pick_first(rows, 'service_title') or '',
        'gst_type': gst_type,
        'gst_percent': str(int(COMPANY_GST_RATE)),
        'taxable': str(_money(taxable)),
        'gst': str(_money(breakdown['gst'])),
        'cgst_amount': str(_money(breakdown['cgst_amount'])),
        'sgst_amount': str(_money(breakdown['sgst_amount'])),
        'igst_amount': str(_money(breakdown['igst_amount'])),
        'total': str(_money(taxable + breakdown['gst'])),
        'lines': lines,
        'source_rows': row_numbers,
    }

    return InvoiceImportItem(
        batch=batch,
        sort_order=sort_order,
        grouping_key=grouping_key[:255],
        invoice_number=invoice_number,
        number_mode=number_mode,
        client=client,
        payload=payload,
        status=status,
        errors=errors,
        warnings=warnings,
    )


def _parse_payload_date(value):
    if not value:
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    try:
        return datetime.strptime(str(value)[:10], '%Y-%m-%d').date()
    except ValueError:
        return None


def _create_invoice_from_item(item, user, batch):
    payload = item.payload or {}
    client = item.client
    if client is None:
        raise ValueError('Cannot create invoice without a customer.')

    invoice_date = _parse_payload_date(payload.get('invoice_date')) or timezone.localdate()
    due_date = _parse_payload_date(payload.get('due_date')) or (invoice_date + timedelta(days=30))
    invoice_number = (item.invoice_number or payload.get('invoice_number') or '').strip()
    number_mode = item.number_mode or Invoice.NUMBER_MODE_MANUAL

    if invoice_number:
        number_mode = Invoice.NUMBER_MODE_MANUAL
    else:
        year_series = get_year_series(invoice_date.year)
        invoice_number = generate_document_number(DOC_INVOICE, user=user, year_series=year_series)
        number_mode = Invoice.NUMBER_MODE_AUTO

    gst_type = payload.get('gst_type') or resolve_client_gst_type(client)
    now = timezone.now()

    invoice = Invoice(
        order=None,
        client=client,
        source=Invoice.SOURCE_IMPORT,
        import_batch=batch,
        invoice_number=invoice_number,
        number_mode=number_mode,
        amount=Decimal('0.00'),
        gst=Decimal('0.00'),
        cgst_amount=Decimal('0.00'),
        sgst_amount=Decimal('0.00'),
        igst_amount=Decimal('0.00'),
        total=Decimal('0.00'),
        gst_type=gst_type,
        invoice_date=invoice_date,
        due_date=due_date,
        po_number=payload.get('po_number') or '',
        po_date=_parse_payload_date(payload.get('po_date')),
        billing_period_from=_parse_payload_date(payload.get('billing_period_from')),
        billing_period_to=_parse_payload_date(payload.get('billing_period_to')),
        service_title=payload.get('service_title') or '',
        approval_status=STATUS_APPROVED,
        created_by=user,
        approved_by=user,
        approved_at=now,
        payment_status='PENDING',
    )
    invoice.save()

    for sl_no, line in enumerate(payload.get('lines') or [], start=1):
        InvoiceLineItem.objects.create(
            invoice=invoice,
            sl_no=sl_no,
            description=line.get('description') or '',
            hsn_sac=line.get('hsn_sac') or COMPANY.get('default_hsn_sac', '998422'),
            unit=line.get('unit') or 'Nos',
            qty=_decimal(line.get('qty')),
            rate=_decimal(line.get('rate')),
            remarks=line.get('remarks') or '',
            from_boq=False,
        )

    invoice.recalculate_totals(gst_type=gst_type)
    invoice.save()

    log_approval_action(
        invoice, ACTION_CREATED, user,
        new_status=STATUS_APPROVED,
        remarks='Created via invoice import.',
    )
    log_approval_action(
        invoice, ACTION_APPROVED, user,
        previous_status='',
        new_status=STATUS_APPROVED,
        remarks='Auto-approved on invoice import.',
    )
    return invoice
