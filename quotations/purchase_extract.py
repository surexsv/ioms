"""Read purchase-rate rows from a spreadsheet, PDF, or photo."""

import csv
import io
import re
import shutil
import subprocess
import zlib
from decimal import Decimal
from pathlib import Path

from .purchase_rates import parse_decimal

ALLOWED_EXTENSIONS = {'.pdf', '.png', '.jpg', '.jpeg', '.webp', '.tif', '.tiff', '.bmp', '.csv', '.xlsx'}
MAX_UPLOAD_BYTES = 8 * 1024 * 1024

_HEADER_ALIASES = {
    'code': ('code', 'item code', 'itemcode', 'material code', 'sku', 'item no', 'item no.'),
    'item': ('item', 'item name', 'name', 'description', 'material', 'product', 'particular', 'particulars'),
    'brand': ('brand', 'make', 'manufacturer'),
    'unit': ('unit', 'uom', 'uom.'),
    'rate': ('rate', 'purchase rate', 'price', 'unit rate', 'cost', 'basic rate', 'basic price', 'purchase price'),
    'gst': ('gst', 'gst%', 'gst %', 'tax', 'tax%', 'tax %', 'gst percent'),
    'remarks': ('remarks', 'remark', 'note', 'notes'),
}
_UNITS = {
    'nos', 'no', 'pcs', 'pc', 'mtr', 'm', 'meter', 'metre', 'kg', 'set', 'ltr', 'box',
    'roll', 'pair', 'ft', 'sqm', 'sqft', 'coil', 'pkt', 'bag', 'lot', 'job', 'day', 'hr', 'hour',
}
_GST = {Decimal('0'), Decimal('5'), Decimal('12'), Decimal('18'), Decimal('28')}
_SKIP_LINE = re.compile(
    r'^(page\b|invoice|tax invoice|gstin|bill to|ship to|total|grand total|sub total|subtotal|amount in words)',
    re.I,
)


class DocumentReadError(Exception):
    pass


def rows_from_upload(uploaded):
    name = Path(uploaded.name or '').name
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentReadError('Upload a PDF, photo (JPG or PNG), CSV, or Excel file.')
    data = uploaded.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentReadError('The file is larger than 8 MB.')
    if not data:
        raise DocumentReadError('The file is empty.')
    if ext == '.csv':
        rows = _rows_from_csv(data)
    elif ext == '.xlsx':
        rows = _rows_from_xlsx(data)
    elif ext == '.pdf':
        rows = parse_rate_text(extract_pdf_text(data))
    else:
        rows = parse_rate_text(extract_image_text(data, ext))
    cleaned = [_clean_row(row) for row in rows]
    cleaned = [row for row in cleaned if row.get('item_name') and row.get('purchase_rate')]
    if not cleaned:
        raise DocumentReadError(
            'No rates were found in that file. Type the rows, or upload a clearer price list with item and rate columns.'
        )
    return cleaned[:200]


def parse_rate_text(text):
    lines = [ln.strip() for ln in (text or '').splitlines() if ln.strip()]
    if not lines:
        return []
    sample = lines[:8]
    if sum(ln.count(',') for ln in sample) >= 2:
        raw = _rows_from_delimited('\n'.join(lines), ',')
    elif any('\t' in ln for ln in sample):
        raw = _rows_from_delimited('\n'.join(lines), '\t')
    else:
        raw = []
        header_at = _header_index(lines)
        if header_at is not None:
            raw = _rows_from_aligned(lines[header_at:])
        if not raw:
            raw = _rows_from_loose(lines)
    cleaned = [_clean_row(row) for row in raw]
    return [row for row in cleaned if row.get('item_name') and row.get('purchase_rate')]


def extract_pdf_text(data):
    text = _extract_pdf_with_pypdf(data)
    if not text.strip():
        text = _extract_pdf_text_fallback(data)
    if not text.strip():
        raise DocumentReadError(
            'This PDF has no readable text. If it is a scan, upload a photo or type the rates.'
        )
    return text


def extract_image_text(data, ext):
    if not shutil.which('tesseract'):
        raise DocumentReadError(
            'Photo reading is not available on this server yet. Upload a PDF or spreadsheet, or type the rates.'
        )
    import tempfile
    suffix = ext if ext else '.png'
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=True) as handle:
        handle.write(data)
        handle.flush()
        try:
            result = subprocess.run(
                ['tesseract', handle.name, 'stdout'],
                capture_output=True,
                timeout=60,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise DocumentReadError('The photo could not be read. Type the rates instead.') from exc
    text = result.stdout.decode('utf-8', errors='ignore')
    if result.returncode != 0 and not text.strip():
        raise DocumentReadError('The photo could not be read. Type the rates instead.')
    return text


def _extract_pdf_with_pypdf(data):
    try:
        from pypdf import PdfReader
    except ImportError:
        return ''
    try:
        reader = PdfReader(io.BytesIO(data))
        return '\n'.join((page.extract_text() or '') for page in reader.pages)
    except Exception:
        return ''


def _extract_pdf_text_fallback(data):
    parts = []
    for match in re.finditer(rb'stream\r?\n(.*?)endstream', data, re.S):
        chunk = _decode_pdf_stream(match.group(1).strip(b'\r\n '))
        parts.append(_strings_from_pdf_content(chunk))
    parts.append(_strings_from_pdf_content(data))
    return '\n'.join(part for part in parts if part)


def _decode_pdf_stream(raw):
    candidates = [raw]
    if raw.endswith(b'~>') or b'~>' in raw:
        try:
            import base64
            candidates.insert(0, base64.a85decode(raw, adobe=True))
        except Exception:
            pass
    for candidate in candidates:
        try:
            return zlib.decompress(candidate)
        except zlib.error:
            continue
    return raw


def _strings_from_pdf_content(chunk):
    try:
        content = chunk.decode('latin1')
    except Exception:
        return ''
    pieces = []
    for array in re.findall(r'\[(.*?)\]\s*TJ', content, re.S):
        pieces.append(_join_pdf_strings(array))
    for literal in re.findall(r'(\((?:\\.|[^\\)])*\)|<([0-9A-Fa-f \n]*)>)\s*Tj', content):
        if literal[0].startswith('('):
            pieces.append(_unescape_pdf_string(literal[0]))
        elif literal[1]:
            pieces.append(_hex_pdf_string(literal[1]))
    return '\n'.join(piece for piece in pieces if piece.strip())


def _join_pdf_strings(array):
    parts = []
    for match in re.finditer(r'\((?:\\.|[^\\)])*\)|<([0-9A-Fa-f \n]*)>', array):
        token = match.group(0)
        if token.startswith('('):
            parts.append(_unescape_pdf_string(token))
        else:
            parts.append(_hex_pdf_string(match.group(1) or ''))
    return ''.join(parts)


def _unescape_pdf_string(token):
    inner = token[1:-1]
    inner = inner.replace(r'\n', '\n').replace(r'\r', '\n').replace(r'\t', '\t')
    inner = inner.replace(r'\(', '(').replace(r'\)', ')').replace(r'\\', '\\')
    return inner


def _hex_pdf_string(hex_text):
    compact = re.sub(r'\s+', '', hex_text)
    if len(compact) % 2:
        compact += '0'
    try:
        return bytes.fromhex(compact).decode('latin1', errors='ignore')
    except ValueError:
        return ''


def _rows_from_csv(data):
    text = data.decode('utf-8-sig', errors='ignore')
    return _rows_from_delimited(text, ',')


def _rows_from_xlsx(data):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    grid = []
    for row in sheet.iter_rows(values_only=True):
        values = ['' if cell is None else str(cell).strip() for cell in row]
        if any(values):
            grid.append(values)
    workbook.close()
    if not grid:
        return []
    headers = [_canonical_header(cell) for cell in grid[0]]
    if not any(headers):
        return _rows_from_loose([' '.join(row) for row in grid])
    rows = []
    for values in grid[1:]:
        mapped = {}
        for key, value in zip(headers, values):
            if key and value:
                mapped[key] = value
        if mapped:
            rows.append(mapped)
    return rows


def _rows_from_delimited(text, delimiter):
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    grid = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]
    if not grid:
        return []
    headers = [_canonical_header(cell) for cell in grid[0]]
    if not any(headers):
        return _rows_from_loose([delimiter.join(row) for row in grid])
    rows = []
    for values in grid[1:]:
        mapped = {}
        for key, value in zip(headers, values):
            if key and value:
                mapped[key] = value
        if mapped:
            rows.append(mapped)
    return rows


def _rows_from_aligned(lines):
    headers = [_canonical_header(cell) for cell in _split_columns(lines[0])]
    if not any(headers):
        return []
    rows = []
    for line in lines[1:]:
        if _SKIP_LINE.search(line):
            continue
        values = _split_columns(line)
        mapped = {}
        for key, value in zip(headers, values):
            if key and value:
                mapped[key] = value
        if mapped:
            rows.append(mapped)
    return rows


def _rows_from_loose(lines):
    rows = []
    for line in lines:
        if _header_index([line]) is not None or _SKIP_LINE.search(line):
            continue
        row = _loose_line(line)
        if row:
            rows.append(row)
    return rows


def _loose_line(line):
    tokens = line.split()
    if len(tokens) < 2:
        return None
    numbers = []
    words = []
    unit = ''
    for token in tokens:
        percent = token.endswith('%')
        bare = token.replace('₹', '').replace('Rs.', '').replace('Rs', '').strip()
        number = parse_decimal(token)
        if number is not None and (percent or _looks_numeric(bare)):
            numbers.append((number, percent))
            continue
        if token.lower().rstrip('.') in _UNITS and not unit:
            unit = token.rstrip('.')
            continue
        words.append(token)
    amounts = []
    gst = None
    for number, percent in numbers:
        if percent and gst is None:
            gst = number
        else:
            amounts.append(number)
    if gst is None and len(amounts) >= 2 and amounts[-1] in _GST and amounts[-1] == amounts[-1].to_integral_value():
        gst = amounts.pop()
    if not amounts or not words:
        return None
    rate = _pick_rate(amounts)
    code = ''
    if words and re.match(r'^[A-Za-z0-9][A-Za-z0-9_./-]{1,20}$', words[0]) and re.search(r'\d', words[0]):
        code = words.pop(0)
    return {
        'code': code,
        'item': ' '.join(words),
        'unit': unit,
        'rate': rate,
        'gst': gst if gst is not None else '',
    }


def _pick_rate(numbers):
    if len(numbers) == 1:
        return numbers[0]
    if len(numbers) >= 3:
        qty, rate, amount = numbers[-3], numbers[-2], numbers[-1]
        if qty > 0 and abs((qty * rate) - amount) <= Decimal('1'):
            return rate
    return numbers[-1]


def _looks_numeric(token):
    return bool(re.match(r'^[\d,]+(\.\d+)?$', token.strip()))


def _header_index(lines):
    for index, line in enumerate(lines[:12]):
        lowered = line.lower()
        if 'rate' in lowered and any(word in lowered for word in ('item', 'description', 'material', 'product', 'particular')):
            return index
    return None


def _split_columns(line):
    if '\t' in line:
        return [cell.strip() for cell in line.split('\t')]
    if line.count(',') >= 2:
        return [cell.strip() for cell in next(csv.reader([line]))]
    return [cell.strip() for cell in re.split(r'\s{2,}', line) if cell.strip()]


def _canonical_header(value):
    label = re.sub(r'\s+', ' ', str(value or '').strip().lower())
    for key, aliases in _HEADER_ALIASES.items():
        if label in aliases:
            return key
    return ''


def _clean_row(row):
    item = row.get('item') or row.get('item_name') or ''
    code = row.get('code') or row.get('item_code') or ''
    rate = parse_decimal(row.get('rate') if row.get('rate') not in (None, '') else row.get('purchase_rate'))
    gst = parse_decimal(row.get('gst') if row.get('gst') not in (None, '') else row.get('gst_percent'))
    return {
        'item_code': str(code or '').strip(),
        'item_name': ' '.join(str(item).split()),
        'brand': str(row.get('brand') or '').strip(),
        'unit': str(row.get('unit') or 'Nos').strip() or 'Nos',
        'purchase_rate': rate,
        'gst_percent': gst if gst is not None else Decimal('18'),
        'remarks': str(row.get('remarks') or '').strip(),
    }
